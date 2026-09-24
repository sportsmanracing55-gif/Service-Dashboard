"""Builds 'Preston Mazda Customer Enquiry Log.xlsx' – a SharePoint-hosted Excel log
that replaces the Node.js tracker. Run: python3 build_workbook.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = "Preston Mazda Customer Enquiry Log.xlsx"
FONT = "Arial"                      # Mazda Type is licensed; the dashboard falls back to Helvetica/Arial
MZ_BLACK, MZ_RED, MZ_SILVER = "101010", "C8102E", "A2AAAD"
INK2, INK3, LINE, SURFACE2 = "4B4F55", "7A7F86", "E2E3E5", "F8F8F8"
GOOD, GOOD_BG, WARN, WARN_BG, INPUT_BG = "1F7A3A", "E3F3E8", "8A5B00", "FFF2C2", "FFFFCC"

ROWS = 200                          # data rows in the table (4..203)
FIRST, LAST = 4, 3 + ROWS

def fill(hex_): return PatternFill("solid", start_color=hex_, end_color=hex_)
def font(size=10, bold=False, color=MZ_BLACK, italic=False): return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
thin = Side(style="thin", color=LINE)
grid = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# =====================================================================
# Enquiry Log
# =====================================================================
ws = wb.active
ws.title = "Enquiry Log"
ws.sheet_view.showGridLines = False

COLS = [  # header, width, group
    ("Ref", 9, ""), ("Date Logged", 12, ""), ("Logged By", 14, ""),
    ("Customer Name", 22, "CUSTOMER"), ("Registration", 12, "CUSTOMER"), ("Contact Number", 15, "CUSTOMER"),
    ("Email Address", 26, "CUSTOMER"), ("Preferred Method", 13, "CUSTOMER"), ("Enquiry", 40, "CUSTOMER"),
    ("Quote File", 24, "PARTS QUOTE (PDF)"), ("Quote Link", 13, "PARTS QUOTE (PDF)"),
    ("Notes", 44, "NOTES"), ("Spoke To Customer", 17, "NOTES"),
    ("Service Advisor", 17, "SERVICE ADVISOR"), ("Contact Customer", 12, "SERVICE ADVISOR"), ("Advisor Done", 11, "SERVICE ADVISOR"),
    ("Parts Person", 17, "PARTS DEPARTMENT"), ("Quote Required", 12, "PARTS DEPARTMENT"), ("Parts Done", 11, "PARTS DEPARTMENT"),
    ("Status", 11, ""),
    ("Parts Alert Sent", 17, "ALERTS – FILLED IN AUTOMATICALLY"), ("Contact Alert Sent", 17, "ALERTS – FILLED IN AUTOMATICALLY"),
    ("Advisor Alerted", 17, "ALERTS – FILLED IN AUTOMATICALLY"), ("Parts Person Alerted", 17, "ALERTS – FILLED IN AUTOMATICALLY"),
    ("Parts Done Alert Sent", 17, "ALERTS – FILLED IN AUTOMATICALLY"),
]
NCOL = len(COLS)
LASTCOL = get_column_letter(NCOL)
col = {h: get_column_letter(i + 1) for i, (h, _, _) in enumerate(COLS)}

# Row 1 – title bar (dashboard top bar: black with red rule)
ws.merge_cells(f"A1:{LASTCOL}1")
ws["A1"] = "PRESTON MAZDA   ·   Customer Enquiry Log"
ws["A1"].font = font(16, True, "FFFFFF"); ws["A1"].fill = fill(MZ_BLACK)
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 34
for c in range(1, NCOL + 1):
    cell = ws.cell(row=1, column=c); cell.fill = fill(MZ_BLACK)
    cell.border = Border(bottom=Side(style="thick", color=MZ_RED))

# Row 2 – column groups
ws.row_dimensions[2].height = 18
i = 0
while i < NCOL:
    g = COLS[i][2]; span = 1
    while i + span < NCOL and COLS[i + span][2] == g and g: span += 1
    c1, c2 = get_column_letter(i + 1), get_column_letter(i + span)
    if span > 1: ws.merge_cells(f"{c1}2:{c2}2")
    cell = ws[f"{c1}2"]; cell.value = g
    cell.font = font(9, True, "FFFFFF" if g else MZ_BLACK); cell.fill = fill(MZ_BLACK if g else "FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    for k in range(i + 1, i + span + 1):
        ws.cell(row=2, column=k).fill = fill(MZ_BLACK if g else "FFFFFF")
        ws.cell(row=2, column=k).border = Border(left=Side(style="thin", color="3A3A40"), right=Side(style="thin", color="3A3A40"))
    i += span

# Row 3 – headers
for idx, (h, w, _) in enumerate(COLS, start=1):
    c = ws.cell(row=3, column=idx, value=h)
    c.font = font(10, True, INK2); c.fill = fill(SURFACE2)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color="C9CBCF"))
    ws.column_dimensions[get_column_letter(idx)].width = w
ws.row_dimensions[3].height = 30

# Data rows: formulas + formats
for r in range(FIRST, LAST + 1):
    ws[f"{col['Ref']}{r}"] = f'=IF({col["Customer Name"]}{r}="","","E-"&TEXT(ROW()-3,"0000"))'
    ws[f"{col['Quote Link']}{r}"] = f'=IF({col["Quote File"]}{r}="","",HYPERLINK(Settings!$B$7&{col["Quote File"]}{r},"Open quote"))'
    ws[f"{col['Status']}{r}"] = (f'=IF({col["Customer Name"]}{r}="","",IF(AND({col["Advisor Done"]}{r}="Yes",'
                                f'OR({col["Quote Required"]}{r}<>"Yes",{col["Parts Done"]}{r}="Yes")),"Complete","Open"))')
    for idx in range(1, NCOL + 1):
        c = ws.cell(row=r, column=idx)
        c.font = font(10); c.border = grid
        c.alignment = Alignment(vertical="top", wrap_text=idx in (9, 12))
    ws[f"{col['Ref']}{r}"].font = font(9, False, INK3); ws[f"{col['Ref']}{r}"].alignment = Alignment(horizontal="center", vertical="top")
    ws[f"{col['Date Logged']}{r}"].number_format = "dd/mm/yyyy"
    ws[f"{col['Customer Name']}{r}"].font = font(10, True)
    ws[f"{col['Registration']}{r}"].font = Font(name="Consolas", size=10)
    ws[f"{col['Quote Link']}{r}"].font = font(10, True, MZ_RED, )
    ws[f"{col['Status']}{r}"].alignment = Alignment(horizontal="center", vertical="top")
    for h in ("Preferred Method", "Contact Customer", "Advisor Done", "Quote Required", "Parts Done"):
        ws[f"{col[h]}{r}"].alignment = Alignment(horizontal="center", vertical="top")
    for h in ("Parts Alert Sent", "Contact Alert Sent", "Advisor Alerted", "Parts Person Alerted", "Parts Done Alert Sent"):
        ws[f"{col[h]}{r}"].font = font(9, False, INK3); ws[f"{col[h]}{r}"].fill = fill(SURFACE2)

# Example row (delete before use – says so in the Setup Guide)
ex = {"Date Logged": "2026-09-24", "Logged By": "Steve Smith", "Customer Name": "Jane Citizen", "Registration": "ABC123",
      "Contact Number": "0400 000 000", "Email Address": "jane@example.com", "Preferred Method": "Text",
      "Enquiry": "Price on front brake pads and rotors for a 2019 CX-5.", "Quote File": "E-0001 Jane Citizen brakes.pdf",
      "Notes": "24/09 Customer called, wants the quote by Friday. – Sam", "Spoke To Customer": "Sam Carter",
      "Service Advisor": "Sam Carter", "Contact Customer": "Yes", "Advisor Done": "No", "Parts Person": "Pat Nguyen",
      "Quote Required": "Yes", "Parts Done": "No"}
from datetime import date
for h, v in ex.items():
    ws[f"{col[h]}{FIRST}"] = date(2026, 9, 24) if h == "Date Logged" else v

# The table (Power Automate reads this by name)
tab = Table(displayName="EnquiryLog", ref=f"A3:{LASTCOL}{LAST}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True, showColumnStripes=False)
ws.add_table(tab)

# Drop-downs
def dv(formula, rng, allow_other=False, prompt=None):
    d = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=not allow_other)
    if allow_other: d.errorStyle = "information"
    if prompt: d.promptTitle, d.prompt, d.showInputMessage = "", prompt, True
    d.error, d.errorTitle = "Please pick a value from the list.", "Not in list"
    ws.add_data_validation(d); d.add(rng)
rng = lambda h: f"{col[h]}{FIRST}:{col[h]}{LAST}"
dv('"Call,Text,Email"', rng("Preferred Method"), prompt="How the customer prefers to be contacted.")
for h in ("Contact Customer", "Advisor Done", "Quote Required", "Parts Done"): dv('"Yes,No"', rng(h))
dv("=Advisors", rng("Service Advisor"), prompt="Advisors come from the Staff sheet.")
dv("=PartsStaff", rng("Parts Person"), prompt="Parts staff come from the Staff sheet.")
dv("=Advisors", rng("Logged By"), allow_other=True)
dv("=Advisors", rng("Spoke To Customer"), allow_other=True)

# Conditional formatting (same meaning as the dashboard: green = done, amber = needs action)
data_rng = f"A{FIRST}:{LASTCOL}{LAST}"
ws.conditional_formatting.add(rng("Advisor Done"), FormulaRule(formula=[f'{col["Advisor Done"]}{FIRST}="Yes"'], fill=fill(GOOD_BG), font=Font(name=FONT, bold=True, color=GOOD)))
ws.conditional_formatting.add(rng("Parts Done"), FormulaRule(formula=[f'{col["Parts Done"]}{FIRST}="Yes"'], fill=fill(GOOD_BG), font=Font(name=FONT, bold=True, color=GOOD)))
ws.conditional_formatting.add(rng("Contact Customer"), FormulaRule(formula=[f'AND({col["Contact Customer"]}{FIRST}="Yes",{col["Advisor Done"]}{FIRST}<>"Yes")'], fill=fill(WARN_BG), font=Font(name=FONT, bold=True, color=WARN)))
ws.conditional_formatting.add(rng("Quote Required"), FormulaRule(formula=[f'AND({col["Quote Required"]}{FIRST}="Yes",{col["Parts Done"]}{FIRST}<>"Yes")'], fill=fill(WARN_BG), font=Font(name=FONT, bold=True, color=WARN)))
ws.conditional_formatting.add(rng("Status"), FormulaRule(formula=[f'{col["Status"]}{FIRST}="Complete"'], fill=fill(GOOD_BG), font=Font(name=FONT, bold=True, color=GOOD)))
ws.conditional_formatting.add(rng("Status"), FormulaRule(formula=[f'{col["Status"]}{FIRST}="Open"'], font=Font(name=FONT, bold=True, color=MZ_RED)))
ws.conditional_formatting.add(f"A{FIRST}:{col['Status']}{LAST}", FormulaRule(formula=[f'${col["Status"]}{FIRST}="Complete"'], fill=fill("F1F8F3")))

ws.freeze_panes = f"{col['Registration']}{FIRST}"   # rows 1-3 and Ref..Customer Name stay put
ws.auto_filter.ref = None
ws.sheet_properties.tabColor = MZ_RED
ws.print_title_rows = "1:3"
ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.print_area = f"A1:{col['Status']}{LAST}"      # the grey alert columns are not needed on paper

# =====================================================================
# Staff
# =====================================================================
st = wb.create_sheet("Staff")
st.sheet_view.showGridLines = False
st.merge_cells("A1:F1"); st["A1"] = "Staff – these names feed the drop-downs; the emails are used by the alert flow"
st["A1"].font = font(12, True, "FFFFFF"); st["A1"].fill = fill(MZ_BLACK); st.row_dimensions[1].height = 26
for c in "ABCDEF": st[f"{c}1"].fill = fill(MZ_BLACK); st[f"{c}1"].border = Border(bottom=Side(style="thick", color=MZ_RED))
st["A2"] = "Service advisors"; st["D2"] = "Parts department"
for c in ("A2", "D2"): st[c].font = font(10, True, INK2)
for c, h in (("A3", "Name"), ("B3", "Email"), ("D3", "Name"), ("E3", "Email")):
    st[c] = h; st[c].font = font(10, True); st[c].fill = fill(SURFACE2); st[c].border = grid
STAFF_LAST = 33
for r in range(4, STAFF_LAST + 1):
    for c in "ABDE":
        st[f"{c}{r}"].border = grid; st[f"{c}{r}"].font = font(10); st[f"{c}{r}"].fill = fill(INPUT_BG)
for r, (n, e) in enumerate([("Sam Carter", "sam@maxkirwan.com.au"), ("Priya Shah", "priya@maxkirwan.com.au")], start=4):
    st[f"A{r}"], st[f"B{r}"] = n, e
st["D4"], st["E4"] = "Pat Nguyen", "pat@maxkirwan.com.au"
st["A35"] = "Yellow cells are for you to fill in. The three names above are examples – replace them with your team. Add more rows as needed (up to row 33)."
st["A35"].font = font(9, False, INK3, italic=True)
for c, w in (("A", 22), ("B", 30), ("C", 3), ("D", 22), ("E", 30)): st.column_dimensions[c].width = w
t1 = Table(displayName="Advisors", ref=f"A3:B{STAFF_LAST}"); t1.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False); st.add_table(t1)
t2 = Table(displayName="PartsStaff", ref=f"D3:E{STAFF_LAST}"); t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False); st.add_table(t2)
wb.defined_names["Advisors"] = DefinedName("Advisors", attr_text=f"Staff!$A$4:$A${STAFF_LAST}")
wb.defined_names["PartsStaff"] = DefinedName("PartsStaff", attr_text=f"Staff!$D$4:$D${STAFF_LAST}")

# =====================================================================
# Settings
# =====================================================================
se = wb.create_sheet("Settings")
se.sheet_view.showGridLines = False
se.merge_cells("A1:C1"); se["A1"] = "Settings"; se["A1"].font = font(12, True, "FFFFFF"); se["A1"].fill = fill(MZ_BLACK); se.row_dimensions[1].height = 26
for c in "ABC": se[f"{c}1"].fill = fill(MZ_BLACK); se[f"{c}1"].border = Border(bottom=Side(style="thick", color=MZ_RED))
rows = [
    ("Dealership", "Preston Mazda", "Shown in alert emails."),
    ("Approver (shares the file)", "steves@maxkirwan.com.au", "Access is approved by sharing the file with a person (SharePoint 'Share'). Nobody else can open it."),
    ("Parts alerts go to", "parts@maxkirwan.com.au", "Emailed when 'Quote Required' is set to Yes."),
    ("Advisor alerts go to", "advisors@maxkirwan.com.au", "Emailed when 'Contact Customer' is set to Yes."),
    ("Quotes folder URL", "https://YOURTENANT.sharepoint.com/sites/Service/Shared%20Documents/Enquiry%20Quotes/", "The SharePoint folder where PDF quotes are saved. Must end with a slash. 'Quote Link' on the log = this + the file name."),
    ("Alert check interval", "5 minutes", "How often the Power Automate flow looks for new alerts to send."),
]
se["A2"], se["B2"], se["C2"] = "Setting", "Value", "Notes"
for c in "ABC": se[f"{c}2"].font = font(10, True); se[f"{c}2"].fill = fill(SURFACE2); se[f"{c}2"].border = grid
for i, (k, v, n) in enumerate(rows, start=3):
    se[f"A{i}"], se[f"B{i}"], se[f"C{i}"] = k, v, n
    se[f"A{i}"].font = font(10, True); se[f"B{i}"].font = font(10, False, "0000FF"); se[f"B{i}"].fill = fill(INPUT_BG); se[f"C{i}"].font = font(9, False, INK3)
    for c in "ABC": se[f"{c}{i}"].border = grid; se[f"{c}{i}"].alignment = Alignment(vertical="top", wrap_text=True)
se.column_dimensions["A"].width = 28; se.column_dimensions["B"].width = 60; se.column_dimensions["C"].width = 70
se["A10"] = "Blue text on yellow = values for you to change."; se["A10"].font = font(9, False, INK3, italic=True)
assert se["A7"].value == "Quotes folder URL"   # Enquiry Log's Quote Link formula points at Settings!B7

# =====================================================================
# Setup Guide
# =====================================================================
g = wb.create_sheet("Setup Guide")
g.sheet_view.showGridLines = False
g.column_dimensions["A"].width = 4; g.column_dimensions["B"].width = 125
g.merge_cells("A1:B1"); g["A1"] = "Setup guide – Preston Mazda Customer Enquiry Log"
g["A1"].font = font(14, True, "FFFFFF"); g["A1"].fill = fill(MZ_BLACK); g["B1"].fill = fill(MZ_BLACK); g.row_dimensions[1].height = 30
for c in "AB": g[f"{c}1"].border = Border(bottom=Side(style="thick", color=MZ_RED))
GUIDE = [
    ("H", "What this is"),
    ("P", "A shared Excel log kept in SharePoint. Everyone with access opens the same file in Excel (desktop or browser) and sees each other's changes within seconds. "
          "Sign-in is your normal Microsoft 365 login, so there are no separate passwords or resets. A Power Automate flow (part of Microsoft 365, runs in the cloud, nothing to install) "
          "sends the email alerts and Teams pop-ups. Set-up takes about 30 minutes and needs no IT installs."),
    ("H", "1. Put the file in SharePoint"),
    ("S", "a) In your Service team's SharePoint site (or a Teams channel's Files tab), create a folder called 'Enquiry Log'."),
    ("S", "b) Upload this workbook into it. Next to it create a folder called 'Enquiry Quotes' – PDF quotes go in there."),
    ("S", "c) Open the 'Enquiry Quotes' folder in the browser, copy the address bar URL, and paste it into Settings → 'Quotes folder URL' (it must end with a slash). "
          "If the URL has '?…' after the folder name, delete that part."),
    ("S", "d) Delete the example row on the Enquiry Log (Jane Citizen) and replace the example names on the Staff sheet with your team and their work email addresses."),
    ("H", "2. Who can use it (approval)"),
    ("S", "Only people the file is shared with can open it. Steve (steves@maxkirwan.com.au) shares it: click 'Share' in SharePoint, type the person's work email, choose 'Can edit'. "
          "To remove someone, 'Manage access' → remove. That is the approval step – nothing else to set up. Forgotten passwords are handled by the normal Microsoft sign-in page."),
    ("H", "3. Day-to-day use"),
    ("S", "• Each customer enquiry is one row. Type in the yellow-free cells; the Ref, Quote Link and Status columns fill themselves."),
    ("S", "• Preferred Method, Service Advisor, Parts Person and the four Yes/No columns are drop-downs (click the cell, then the arrow). Contact Customer = Yes means a service advisor must ring/text/email the customer. Quote Required = Yes means the parts department must prepare a quote."),
    ("S", "• Each department marks its own task: Advisor Done and Parts Done. When both are done (or no quote was needed) Status shows 'Complete' and the row turns green. Use the filter arrow on Status to hide complete rows."),
    ("S", "• Attaching a quote: save the PDF into the 'Enquiry Quotes' folder named with the Ref, e.g. 'E-0012 Smith brakes.pdf', then type that file name into Quote File. 'Open quote' becomes a link."),
    ("S", "• Notes: write what was discussed in Notes (Alt+Enter for a new line, start each note with the date) and put the name of the person who spoke to the customer in Spoke To Customer. "
          "For a running conversation you can also right-click the row → New Comment; comments record who wrote them and when."),
    ("S", "• Date Logged: press Ctrl+; to insert today's date. Logged By: your name."),
    ("S", "• Do not sort the table (use the filter arrows instead) – the Ref numbers follow the row position. Filtering is fine."),
    ("S", "• The grey 'ALERTS' columns on the right are written by the flow. Leave them alone; clearing a cell makes the flow send that alert again."),
    ("H", "4. Email alerts and pop-ups (Power Automate) – build once, about 20 minutes"),
    ("P", "Go to https://make.powerautomate.com, sign in with your work account, click 'Create' → 'Scheduled cloud flow'. Name: 'Enquiry log alerts'. Repeat every 5 minutes. Then add these steps in order:"),
    ("S", "STEP 1 – 'List rows present in a table' (Excel Online (Business)). Location: your SharePoint site · Document Library: Documents · File: Enquiry Log/…xlsx · Table: EnquiryLog. "
          "Under advanced options set Filter Query to:  Customer Name ne ''  (so blank rows are skipped)."),
    ("S", "STEP 2 – 'List rows present in a table' again for Table: Advisors (same file), and a third one for Table: PartsStaff. Rename them 'Advisors' and 'Parts staff'."),
    ("S", "STEP 3 – 'Apply to each' over the output of Step 1 ('value'). Inside it add the following Conditions (each one is a separate Condition block, one under the other, not nested):"),
    ("S", "  A) PARTS QUOTE REQUIRED  →  Condition: 'Quote Required' is equal to Yes  AND  'Parts Alert Sent' is equal to '' (empty). "
          "If yes: 'Send an email (V2)' To: parts@maxkirwan.com.au, Subject: Parts quote required – [Customer Name] [Registration], Body: the row's details (Ref, name, rego, phone, email, preferred method, enquiry, service advisor). "
          "Then 'Update a row' (Excel Online): same file/table, Key Column: Ref, Key Value: [Ref], and set Parts Alert Sent to the expression  formatDateTime(utcNow(),'dd/MM/yyyy HH:mm')."),
    ("S", "  B) ADVISOR MUST CONTACT CUSTOMER  →  Condition: 'Contact Customer' is equal to Yes AND 'Contact Alert Sent' is equal to ''. "
          "If yes: Send an email (V2) To: advisors@maxkirwan.com.au, Subject: Customer contact required – [Customer Name] [Registration], Body: details + 'Preferred method: [Preferred Method]'. Then Update a row → Contact Alert Sent = formatDateTime(utcNow(),'dd/MM/yyyy HH:mm')."),
    ("S", "  C) ALLOCATED TO A SERVICE ADVISOR (email + pop-up to that person)  →  Condition: 'Service Advisor' is not equal to '' AND 'Service Advisor' is not equal to 'Advisor Alerted'. "
          "If yes: add a 'Filter array' on the Advisors list (Step 2) where Name is equal to [Service Advisor]; then Send an email (V2) To: the expression  first(body('Filter_array'))?['Email']  , "
          "Subject: Enquiry [Ref] allocated to you – [Customer Name], Body: details. Then 'Post message in a chat or channel' (Microsoft Teams): Post as Flow bot, Post in Chat with Flow bot, Recipient: the same email expression, "
          "Message: 'Enquiry [Ref] ([Customer Name], [Registration]) has been allocated to you. [Enquiry]'. This appears as a Teams pop-up notification on their PC and phone. Then Update a row → Advisor Alerted = [Service Advisor]."),
    ("S", "  D) ALLOCATED TO A PARTS PERSON  →  same as C using 'Parts Person', the Parts staff list, and Update a row → Parts Person Alerted = [Parts Person]."),
    ("S", "  E) PARTS QUOTE DONE – tell the advisor  →  Condition: 'Parts Done' is equal to Yes AND 'Parts Done Alert Sent' is equal to '' AND 'Service Advisor' is not equal to ''. "
          "If yes: look up the advisor's email as in C, Send an email (V2) + Teams message 'Parts quote ready for [Customer Name] ([Ref])', then Update a row → Parts Done Alert Sent = formatDateTime(utcNow(),'dd/MM/yyyy HH:mm')."),
    ("S", "STEP 4 – Save, then click 'Test' → Manually → Run. Set a row's Quote Required to Yes, wait 5 minutes and check parts@ received the email and the grey 'Parts Alert Sent' cell filled in."),
    ("P", "Tips: in a Condition, pick the column from 'Dynamic content' (they are listed under the Step 1 action). For 'is equal to' blank, leave the right-hand box empty. "
          "If you converted the Yes/No columns to Excel check boxes (Insert → Checkbox), compare with 'true' instead of 'Yes'. "
          "If Power Automate says it is not available to you, ask IT to enable it for your account – it is included in Microsoft 365 business plans. Until the flow exists the log still works; only the automatic alerts are missing."),
    ("H", "5. Where the data lives"),
    ("P", "Everything is in this workbook and the Quotes folder on your SharePoint site – inside the dealership's own Microsoft 365 tenant, under IT's existing controls and backups. Nothing is on any other website. "
          "SharePoint keeps version history: right-click the file → Version history to restore an earlier copy if something is deleted by mistake."),
]
r = 3
for kind, text in GUIDE:
    c = g[f"B{r}"]; c.value = text
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "H":
        c.font = font(12, True, MZ_RED); g.row_dimensions[r].height = 22
    else:
        c.font = font(10)
        g.row_dimensions[r].height = max(16, 15 * (len(text) // 120 + 1))
    r += 1
g.sheet_properties.tabColor = MZ_SILVER
g.page_setup.orientation = "landscape"; g.page_setup.paperSize = g.PAPERSIZE_A4
g.sheet_properties.pageSetUpPr.fitToPage = True; g.page_setup.fitToWidth = 1; g.page_setup.fitToHeight = 0
g.page_margins.left = g.page_margins.right = 0.4

wb.move_sheet("Setup Guide", offset=-3)   # first tab people see
wb.active = 1                               # but open on the log
wb.save(OUT)
print("wrote", OUT)
